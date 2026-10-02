"""Excel BOQ formatter — outputs a professional styled .xlsx workbook."""
from __future__ import annotations
from pathlib import Path
from typing import Any

def format_excel(boq: Any, output_path: str | Path) -> Path:
    """Write *boq* to a styled Excel workbook.

    Returns the saved path. Requires openpyxl.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl required: pip install openpyxl")

    p = Path(output_path)
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ"

    # Styles
    hdr_font = Font(bold=True, color="FFFFFF", size=11, name="Arial")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    sec_fill = PatternFill("solid", fgColor="BDD7EE")
    alt_fill = PatternFill("solid", fgColor="F2F2F2")
    cen      = Alignment(horizontal="center", vertical="center")
    right    = Alignment(horizontal="right")
    thin     = Side(style="thin")
    bdr      = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _cell(r, c, val, bold=False, fill=None, align=None, num_fmt=None):
        cell = ws.cell(row=r, column=c, value=val)
        if bold:  cell.font = Font(bold=True, name="Arial")
        if fill:  cell.fill = fill
        if align: cell.alignment = align
        if num_fmt: cell.number_format = num_fmt
        cell.border = bdr
        return cell

    # Title rows
    ws.merge_cells("A1:I1")
    t = ws.cell(row=1, column=1, value=f"ABSTRACT BILL OF QUANTITIES — {boq.project_name}")
    t.font = Font(bold=True, size=14, color="FFFFFF", name="Arial")
    t.fill = PatternFill("solid", fgColor="1F3864")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:I2")
    ws.cell(row=2, column=1,
            value=f"Bridge: {boq.bridge_name}  |  Chainage: {boq.chainage_km} km")
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    # Headers
    headers = ["Item", "Description", "Unit", "Quantity",
               "Rate (INR)", "Amount (INR)", "MORTH Ref", "Remarks", ""]
    col_w   = [10, 60, 10, 14, 16, 18, 14, 20, 5]
    for c, (h, w) in enumerate(zip(headers, col_w), 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = cen
        cell.border = bdr
        ws.column_dimensions[get_column_letter(c)].width = w

    row = 4
    for sec in boq.sections:
        # Section header
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        sh = ws.cell(row=row, column=1,
                     value=f"{sec.section_code}  {sec.title}")
        sh.font = Font(bold=True, name="Arial", size=10)
        sh.fill = sec_fill
        sh.border = bdr
        ws.row_dimensions[row].height = 16
        row += 1

        for i, item in enumerate(sec.items):
            fill = alt_fill if i % 2 else None
            _cell(row, 1, item.item_code,   fill=fill, align=cen)
            _cell(row, 2, item.description, fill=fill)
            _cell(row, 3, item.unit,        fill=fill, align=cen)
            _cell(row, 4, float(item.quantity), fill=fill, align=right,
                  num_fmt="#,##0.00")
            _cell(row, 5, float(item.rate) if item.rate else "TBC",
                  fill=fill, align=right, num_fmt="#,##0.00")
            _cell(row, 6, float(item.amount) if item.amount else 0,
                  fill=fill, align=right, num_fmt="₹#,##0.00")
            _cell(row, 7, item.ref,     fill=fill, align=cen)
            _cell(row, 8, item.remarks, fill=fill)
            row += 1

        # Section total
        st = ws.cell(row=row, column=1,
                     value=f"Total — {sec.title}")
        st.font = Font(bold=True, name="Arial")
        st.border = bdr
        for c in range(2, 9):
            ws.cell(row=row, column=c).border = bdr
        amt = ws.cell(row=row, column=6,
                      value=float(sec.total))
        amt.font = Font(bold=True, name="Arial")
        amt.number_format = "₹#,##0.00"
        amt.alignment = right
        amt.border = bdr
        row += 2

    # Grand total rows
    for label, val in [
        ("SUB-TOTAL (Civil Works)",      boq.subtotal),
        (f"CONTINGENCY @ {boq.contingency_pct}%", boq.contingency_amount),
        ("GRAND TOTAL",                   boq.grand_total),
    ]:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=11, name="Arial")
        c = ws.cell(row=row, column=6, value=float(val))
        c.font = Font(bold=True, size=11, name="Arial")
        c.number_format = "₹#,##0.00"
        c.alignment = right
        for col in range(1, 9):
            ws.cell(row=row, column=col).border = bdr
        row += 1

    wb.save(str(p))
    return p


__all__ = ["format_excel"]
