"""
bridgecad_io.excel_reader — Parse a 14-sheet BridgeCAD Excel workbook into a
validated BridgeProject Pydantic model.

Sheet order (1-indexed):
  S01 PROJECT_MASTER      S02 BRIDGE_SELECTION    S03 GEOMETRY_INPUT
  S04 SUPERSTRUCTURE      S05 SUBSTRUCTURE        S06 FOUNDATION_DETAILS
  S07 APPROACHES          S08 HYDRAULIC_DATA       S09 MATERIALS
  S10 BEARINGS_JOINTS     S11 COMPONENTS_LIB       S12 DRAWING_CONTROL
  S13 CALCULATIONS        S14 VALIDATION

Each sheet uses the canonical 3-column layout:
  Col A: Field name (matches Pydantic field names, case-insensitive)
  Col B: Value
  Col C: Unit / Notes (ignored during parse)

Bearing and expansion-joint schedules (S10) use a tabular block starting
at the row after the ``bearing_schedule:`` / ``expansion_joint_schedule:``
sentinel rows, with column headers in the next row.

M2 Week 2 implementation.
"""
from __future__ import annotations

import logging
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

try:
    import openpyxl
    _OPENPYXL_AVAILABLE = True
except ImportError:
    _OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl not installed — excel_reader will not function")

# Lazy import of bridgecad_core so bridgecad_io can be imported even when
# bridgecad_core is not yet installed (test environments, CI stub runs).
def _core():
    from bridgecad_core import models as _m, types as _t
    return _m, _t


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _cell_value(ws, row: int, col: int) -> Any:
    """Return the value of a cell, stripping leading/trailing whitespace."""
    v = ws.cell(row=row, column=col).value
    if isinstance(v, str):
        v = v.strip()
    return v


def _sheet_to_dict(ws) -> dict[str, Any]:
    """Read a 3-column (Field / Value / Notes) sheet into a plain dict."""
    result: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip().lower().replace(" ", "_")
        value = row[1] if len(row) > 1 else None
        if key and value is not None and value != "":
            result[key] = value
    return result


def _coerce(value: Any, python_type: type) -> Any:
    """Best-effort coercion from Excel cell value to target Python type."""
    if value is None:
        return None
    if python_type is bool:
        if isinstance(value, bool):
            return value
        return str(value).strip().upper() in {"TRUE", "YES", "1", "Y"}
    if python_type is int:
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return None
    if python_type is float:
        try:
            return float(value)
        except (TypeError, ValueError):
            return None
    if python_type is Decimal:
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    if python_type is date:
        if isinstance(value, date):
            return value
        try:
            from datetime import datetime
            return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
        except (ValueError, TypeError):
            return None
    if python_type is str:
        return str(value).strip() if value is not None else None
    return value


def _lookup_enum(enum_cls, raw: Any) -> Any:
    """Find an enum member by name or value; return None if not found."""
    if raw is None:
        return None
    s = str(raw).strip()
    # Try by name first (case-insensitive)
    for m in enum_cls:
        if m.name.upper() == s.upper():
            return m
    # Try by value
    for m in enum_cls:
        if str(m.value).upper() == s.upper():
            return m
    logger.warning("Could not resolve %s member from %r", enum_cls.__name__, raw)
    return None


# ---------------------------------------------------------------------------
# Per-sheet parsers
# ---------------------------------------------------------------------------

def _parse_project_master(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "project_title":            _coerce(d.get("project_title"), str),
        "project_code":             _coerce(d.get("project_code"), str),
        "bridge_name":              _coerce(d.get("bridge_name"), str),
        "chainage_km":              _coerce(d.get("chainage_km"), Decimal),
        "package_no":               _coerce(d.get("package_no"), str),
        "client":                   _lookup_enum(t.ClientType, d.get("client")),
        "consultant":               _lookup_enum(t.ConsultantType, d.get("consultant")),
        "contractor":               _lookup_enum(t.ContractorType, d.get("contractor")),
        "drawing_no":               _coerce(d.get("drawing_no"), str),
        "drawing_title":            _coerce(d.get("drawing_title"), str),
        "revision":                 _lookup_enum(t.RevisionTag, d.get("revision")),
        "date_issued":              _coerce(d.get("date_issued"), date),
        "designed_by":              _coerce(d.get("designed_by"), str),
        "checked_by":               _coerce(d.get("checked_by"), str),
        "approved_by":              _coerce(d.get("approved_by"), str),
        "road_level_rl_m":          _coerce(d.get("road_level_rl_m"), Decimal),
        "ground_level_rl_m":        _coerce(d.get("ground_level_rl_m"), Decimal),
        "survey_date":              _coerce(d.get("survey_date"), date),
        "state_code":               _lookup_enum(t.IndianStateCode, d.get("state_code")),
        "district_code":            _coerce(d.get("district_code"), str),
        "language":                 _lookup_enum(t.Language, d.get("language")),
        "latitude_deg":             _coerce(d.get("latitude_deg"), Decimal),
        "longitude_deg":            _coerce(d.get("longitude_deg"), Decimal),
        "total_estimated_cost_inr": _coerce(d.get("total_estimated_cost_inr"), Decimal),
        "tender_no":                _coerce(d.get("tender_no"), str),
        "contract_no":              _coerce(d.get("contract_no"), str),
        "project_phase":            _lookup_enum(t.ProjectPhase, d.get("project_phase")),
        "currency":                 _lookup_enum(t.Currency, d.get("currency")),
    }


def _parse_bridge_selection(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "bridge_category":    _lookup_enum(t.BridgeCategory, d.get("bridge_category")),
        "bridge_type":        _lookup_enum(t.SuperstructureType, d.get("bridge_type")),
        "span_configuration": _lookup_enum(t.SpanConfiguration, d.get("span_configuration")),
        "span_count":         _coerce(d.get("span_count"), int),
        "carriageway_config": _lookup_enum(t.CarriagewayLaneConfig, d.get("carriageway_config")),
        "design_code":        _lookup_enum(t.DesignCode, d.get("design_code")),
        "loading_class":      _lookup_enum(t.LoadClass, d.get("loading_class")),
        "seismic_zone":       _lookup_enum(t.SeismicZone, d.get("seismic_zone")),
        "wind_zone_ms":       _lookup_enum(t.WindSpeedBasic_ms, d.get("wind_zone_ms")),
        "foundation_type":    _lookup_enum(t.FoundationType, d.get("foundation_type")),
    }


def _parse_geometry_input(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    # span_lengths_m may be comma-separated or individual keys span_length_1 .. N
    spans_raw = d.get("span_lengths_m", "")
    span_lengths: list[Decimal] = []
    if spans_raw:
        for part in str(spans_raw).split(","):
            v = _coerce(part.strip(), Decimal)
            if v is not None:
                span_lengths.append(v)
    # Fallback: span_length_1, span_length_2, ...
    if not span_lengths:
        i = 1
        while True:
            key = f"span_length_{i}"
            if key not in d:
                break
            v = _coerce(d[key], Decimal)
            if v is not None:
                span_lengths.append(v)
            i += 1
    span_count = _coerce(d.get("span_count"), int) or len(span_lengths) or 1
    if not span_lengths:
        span_lengths = [Decimal("12.0")]
    return {
        "span_count":                       span_count,
        "span_lengths_m":                   span_lengths,
        "alignment_type":                   _lookup_enum(t.AlignmentType, d.get("alignment_type")),
        "gradient_pct":                     _coerce(d.get("gradient_pct"), Decimal),
        "carriageway_width_m":              _coerce(d.get("carriageway_width_m"), Decimal),
        "footpath_left_m":                  _coerce(d.get("footpath_left_m"), Decimal),
        "footpath_right_m":                 _coerce(d.get("footpath_right_m"), Decimal),
        "crash_barrier_width_left_m":       _coerce(d.get("crash_barrier_width_left_m"), Decimal),
        "crash_barrier_width_right_m":      _coerce(d.get("crash_barrier_width_right_m"), Decimal),
        "kerb_width_left_m":                _coerce(d.get("kerb_width_left_m"), Decimal),
        "kerb_width_right_m":               _coerce(d.get("kerb_width_right_m"), Decimal),
        "median_width_m":                   _coerce(d.get("median_width_m"), Decimal),
        "skew_angle_deg":                   _coerce(d.get("skew_angle_deg"), Decimal),
        "skew_direction":                   _lookup_enum(t.SkewDirection, d.get("skew_direction")),
        "curved_bridge_flag":               _coerce(d.get("curved_bridge_flag"), bool),
        "camber_mm":                        _coerce(d.get("camber_mm"), int),
        "camber_method":                    _lookup_enum(t.CamberMethod, d.get("camber_method")),
        "pavement_thickness_total_mm":      _coerce(d.get("pavement_thickness_total_mm"), int),
        "crust_layer_count":                _coerce(d.get("crust_layer_count"), int),
        "chainage_unit":                    _lookup_enum(t.ChainageUnit, d.get("chainage_unit")),
        "chainage_start_km":                _coerce(d.get("chainage_start_km"), Decimal),
        "chainage_end_km":                  _coerce(d.get("chainage_end_km"), Decimal),
        "bearing_pad_spacing_m":            _coerce(d.get("bearing_pad_spacing_m"), Decimal),
        "kerb_type":                        _lookup_enum(t.KerbType, d.get("kerb_type")),
        "footpath_type":                    _lookup_enum(t.FootpathType, d.get("footpath_type")),
        "median_type":                      _lookup_enum(t.MedianType, d.get("median_type")),
        "crash_barrier_type":               _lookup_enum(t.CrashBarrierType, d.get("crash_barrier_type")),
    }


def _parse_superstructure(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "type":                            _lookup_enum(t.SuperstructureType, d.get("type")),
        "deck_material":                   _lookup_enum(t.ConcreteGrade, d.get("deck_material")),
        "deck_thickness_mm":               _coerce(d.get("deck_thickness_mm"), int),
        "wearing_coat_type":               _lookup_enum(t.WearingCoatType, d.get("wearing_coat_type")),
        "wearing_coat_grade":              _lookup_enum(t.WearingCoatGrade, d.get("wearing_coat_grade")),
        "wearing_coat_thickness_mm":       _coerce(d.get("wearing_coat_thickness_mm"), int),
        "wearing_coat_slope_pct":          _coerce(d.get("wearing_coat_slope_pct"), Decimal),
        "drainage_type":                   _lookup_enum(t.DrainageType, d.get("drainage_type")),
        "drainage_spacing_m":              _coerce(d.get("drainage_spacing_m"), Decimal),
        "camber_mm":                       _coerce(d.get("camber_mm"), int),
        "camber_method":                   _lookup_enum(t.CamberMethod, d.get("camber_method")),
        "girder_type":                     _lookup_enum(t.GirderType, d.get("girder_type")),
        "girder_depth_mm":                 _coerce(d.get("girder_depth_mm"), int),
        "girder_spacing_m":                _coerce(d.get("girder_spacing_m"), Decimal),
        "girders_per_deck_count":          _coerce(d.get("girders_per_deck_count"), int),
        "parapet_type":                    _lookup_enum(t.ParapetType, d.get("parapet_type")),
        "parapet_height_mm":               _coerce(d.get("parapet_height_mm"), int),
        "bearing_shelf_width_m":           _coerce(d.get("bearing_shelf_width_m"), Decimal),
        "transverse_diaphragm_spacing_m":  _coerce(d.get("transverse_diaphragm_spacing_m"), Decimal),
        "continuity_slab_thickness_mm":    _coerce(d.get("continuity_slab_thickness_mm"), int),
        "kerb_elevation_mm":               _coerce(d.get("kerb_elevation_mm"), int),
        "footpath_elevation_mm":           _coerce(d.get("footpath_elevation_mm"), int),
    }


def _parse_substructure(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "pier_type":                  _lookup_enum(t.PierType, d.get("pier_type")),
        "pier_material":              _lookup_enum(t.ConcreteGrade, d.get("pier_material")),
        "pier_height_typical_m":      _coerce(d.get("pier_height_typical_m"), Decimal),
        "pier_cap_type":              _lookup_enum(t.PierCapType, d.get("pier_cap_type")),
        "pier_cap_width_m":           _coerce(d.get("pier_cap_width_m"), Decimal),
        "pier_cap_length_m":          _coerce(d.get("pier_cap_length_m"), Decimal),
        "pier_cap_depth_m":           _coerce(d.get("pier_cap_depth_m"), Decimal),
        "pier_shaft_width_m":         _coerce(d.get("pier_shaft_width_m"), Decimal),
        "pier_shaft_length_m":        _coerce(d.get("pier_shaft_length_m"), Decimal),
        "pier_pedestal_height_m":     _coerce(d.get("pier_pedestal_height_m"), Decimal),
        "pier_diaphragm_required":    _coerce(d.get("pier_diaphragm_required"), bool),
        "abutment_type":              _lookup_enum(t.AbutmentType, d.get("abutment_type")),
        "abutment_material":          _lookup_enum(t.ConcreteGrade, d.get("abutment_material")),
        "abutment_width_m":           _coerce(d.get("abutment_width_m"), Decimal),
        "abutment_length_m":          _coerce(d.get("abutment_length_m"), Decimal),
        "abutment_height_m":          _coerce(d.get("abutment_height_m"), Decimal),
        "abutment_pedestal_height_m": _coerce(d.get("abutment_pedestal_height_m"), Decimal),
        "return_wall_type":           _lookup_enum(t.ReturnWallType, d.get("return_wall_type")),
        "return_wall_length_m":       _coerce(d.get("return_wall_length_m"), Decimal),
        "return_wall_height_m":       _coerce(d.get("return_wall_height_m"), Decimal),
        "wing_wall_type":             _lookup_enum(t.WingWallType, d.get("wing_wall_type")),
        "wing_wall_angle_deg":        _coerce(d.get("wing_wall_angle_deg"), Decimal),
        "wing_wall_splay_length_m":   _coerce(d.get("wing_wall_splay_length_m"), Decimal),
        "wing_wall_height_m":         _coerce(d.get("wing_wall_height_m"), Decimal),
        "backfill_type":              _coerce(d.get("backfill_type"), str),
        "approach_slab_length_m":     _coerce(d.get("approach_slab_length_m"), Decimal),
        "approach_slab_thickness_mm": _coerce(d.get("approach_slab_thickness_mm"), int),
        "bearings_per_pier":          _coerce(d.get("bearings_per_pier"), int),
        "bearings_per_abutment":      _coerce(d.get("bearings_per_abutment"), int),
    }


def _parse_foundation_details(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "type":                          _lookup_enum(t.FoundationType, d.get("type")),
        "pile_type":                     _lookup_enum(t.PileType, d.get("pile_type")),
        "pile_shape":                    _lookup_enum(t.PileShape, d.get("pile_shape")),
        "pile_diameter_mm":              _coerce(d.get("pile_diameter_mm"), int),
        "pile_length_m":                 _coerce(d.get("pile_length_m"), Decimal),
        "piles_per_pier":                _coerce(d.get("piles_per_pier"), int),
        "pile_group_config":             _coerce(d.get("pile_group_config"), str),
        "pile_spacing_m":                _coerce(d.get("pile_spacing_m"), Decimal),
        "pile_cap_thickness_mm":         _coerce(d.get("pile_cap_thickness_mm"), int),
        "pile_cutoff_level_m":           _coerce(d.get("pile_cutoff_level_m"), Decimal),
        "under_ream_count":              _lookup_enum(t.UnderReamCount, d.get("under_ream_count")),
        "factor_of_safety_on_bearing":   _coerce(d.get("factor_of_safety_on_bearing"), Decimal),
        "bearing_capacity_kpa":          _coerce(d.get("bearing_capacity_kpa"), Decimal),
    }


def _parse_approaches(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "left_approach_length_m":              _coerce(d.get("left_approach_length_m"), Decimal),
        "right_approach_length_m":             _coerce(d.get("right_approach_length_m"), Decimal),
        "type_of_embankment":                  _coerce(d.get("type_of_embankment"), str),
        "soil_embankment_unit_weight_kn_m3":   _coerce(d.get("soil_embankment_unit_weight_kn_m3"), Decimal),
        "approach_slope_width_m":              _coerce(d.get("approach_slope_width_m"), Decimal),
        "shoulder_width_left_m":               _coerce(d.get("shoulder_width_left_m"), Decimal),
        "shoulder_width_right_m":              _coerce(d.get("shoulder_width_right_m"), Decimal),
        "guide_rail_required":                 _coerce(d.get("guide_rail_required"), bool),
        "speed_limit_on_approach_kmph":        _coerce(d.get("speed_limit_on_approach_kmph"), int),
        "transition_curve_length_m":           _coerce(d.get("transition_curve_length_m"), Decimal),
        "approach_gradient_pct":               _coerce(d.get("approach_gradient_pct"), Decimal),
    }


def _parse_hydraulic_data(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "design_discharge_cumecs":              _coerce(d.get("design_discharge_cumecs"), Decimal),
        "hfl_m":                                _coerce(d.get("hfl_m"), Decimal),
        "lwl_m":                                _coerce(d.get("lwl_m"), Decimal),
        "normal_wl_m":                          _coerce(d.get("normal_wl_m"), Decimal),
        "scour_level_m":                        _coerce(d.get("scour_level_m"), Decimal),
        "design_scour_depth_m":                 _coerce(d.get("design_scour_depth_m"), Decimal),
        "freeboard_m":                          _coerce(d.get("freeboard_m"), Decimal),
        "actual_soffit_level_m":                _coerce(d.get("actual_soffit_level_m"), Decimal),
        "waterway_required_m2":                 _coerce(d.get("waterway_required_m2"), Decimal),
        "waterway_provided_m2":                 _coerce(d.get("waterway_provided_m2"), Decimal),
        "lacey_silt_factor":                    _coerce(d.get("lacey_silt_factor"), Decimal),
        "regime_perimeter_m":                   _coerce(d.get("regime_perimeter_m"), Decimal),
        "regime_depth_m":                       _coerce(d.get("regime_depth_m"), Decimal),
        "regime_velocity_mps":                  _coerce(d.get("regime_velocity_mps"), Decimal),
        "afflux_m":                             _coerce(d.get("afflux_m"), Decimal),
        "max_afflux_allowable_m":               _coerce(d.get("max_afflux_allowable_m"), Decimal),
        "water_source_type":                    _lookup_enum(t.WaterSourceType, d.get("water_source_type")),
        "river_bed_slope_m_per_km":             _coerce(d.get("river_bed_slope_m_per_km"), Decimal),
        "catchment_area_sqkm":                  _coerce(d.get("catchment_area_sqkm"), Decimal),
        "return_period_yr":                     _lookup_enum(t.FloodReturnPeriodDesign, d.get("return_period_yr")),
        "any_regime_equation_not_applied_flag": _coerce(d.get("any_regime_equation_not_applied_flag"), bool),
        "highest_recorded_flood_year":          _coerce(d.get("highest_recorded_flood_year"), int),
    }


def _parse_materials(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "concrete_superstructure":     _lookup_enum(t.ConcreteGrade, d.get("concrete_superstructure")),
        "concrete_substructure":       _lookup_enum(t.ConcreteGrade, d.get("concrete_substructure")),
        "concrete_foundation":         _lookup_enum(t.ConcreteGrade, d.get("concrete_foundation")),
        "reinforcement_steel_grade":   _lookup_enum(t.SteelGrade, d.get("reinforcement_steel_grade")),
        "prestressing_steel":          _lookup_enum(t.PrestressGrade, d.get("prestressing_steel")),
        "bearing_type":                _lookup_enum(t.BearingType, d.get("bearing_type")),
        "expansion_joint_type":        _lookup_enum(t.ExpansionJointType, d.get("expansion_joint_type")),
        "wearing_coat_material_type":  _lookup_enum(t.WearingCoatType, d.get("wearing_coat_material_type")),
        "wearing_coat_grade":          _lookup_enum(t.WearingCoatGrade, d.get("wearing_coat_grade")),
    }


def _parse_bearings_joints(ws, m, t) -> dict:
    """Parse S10 — bearing_schedule and expansion_joint_schedule tabular blocks."""
    bearing_rows: list[dict] = []
    expansion_rows: list[dict] = []

    rows = list(ws.iter_rows(values_only=True))
    mode: str | None = None
    headers: list[str] = []

    for row in rows:
        if not row or row[0] is None:
            continue
        key0 = str(row[0]).strip().lower()
        if key0 in ("bearing_schedule:", "bearing_schedule"):
            mode = "bearing"; headers = []; continue
        if key0 in ("expansion_joint_schedule:", "expansion_joint_schedule"):
            mode = "expansion"; headers = []; continue
        if mode and not headers:
            # Next non-empty row after sentinel = header row
            headers = [str(c).strip().lower().replace(" ", "_") if c else "" for c in row]
            continue
        if mode and headers:
            if all(c is None or str(c).strip() == "" for c in row):
                mode = None; headers = []; continue
            rec = {}
            for i, h in enumerate(headers):
                if h:
                    rec[h] = row[i] if i < len(row) else None
            if any(v is not None and str(v).strip() != "" for v in rec.values()):
                if mode == "bearing":
                    bearing_rows.append(rec)
                else:
                    expansion_rows.append(rec)

    def _mk_bearing(r: dict) -> dict:
        return {
            "location":               _coerce(r.get("location"), str),
            "bearing_type":           _lookup_enum(t.BearingType, r.get("bearing_type")),
            "size_x_mm":              _coerce(r.get("size_x_mm"), int),
            "size_y_mm":              _coerce(r.get("size_y_mm"), int),
            "design_load_kn":         _coerce(r.get("design_load_kn"), Decimal),
            "fixed_or_guide_or_free": _coerce(r.get("fixed_or_guide_or_free"), str),
            "elastomeric_layers_count": _coerce(r.get("elastomeric_layers_count"), int),
            "steel_back_bool":        _coerce(r.get("steel_back_bool"), bool),
            "pot_pressure_mpa":       _coerce(r.get("pot_pressure_mpa"), Decimal),
            "ptfe_sliding_surface_flag": _coerce(r.get("ptfe_sliding_surface_flag"), bool),
            "quantity":               _coerce(r.get("quantity"), int) or 1,
        }

    def _mk_expansion(r: dict) -> dict:
        return {
            "location":           _coerce(r.get("location"), str),
            "joint_type":         _lookup_enum(t.ExpansionJointType, r.get("joint_type")),
            "width_mm":           _coerce(r.get("width_mm"), int),
            "movement_range_mm":  _coerce(r.get("movement_range_mm"), int),
            "gap_mm":             _coerce(r.get("gap_mm"), int),
            "transverse_length_m": _coerce(r.get("transverse_length_m"), Decimal),
            "quantity":           _coerce(r.get("quantity"), int) or 1,
        }

    return {
        "bearing_schedule":         [_mk_bearing(r) for r in bearing_rows],
        "expansion_joint_schedule": [_mk_expansion(r) for r in expansion_rows],
    }


def _parse_components_lib(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {k: _coerce(v, str) for k, v in d.items()}


def _parse_drawing_control(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    # output_formats may be comma-separated
    fmt_raw = d.get("output_formats", "")
    output_formats = []
    for f in str(fmt_raw).split(","):
        e = _lookup_enum(t.OutputFormat, f.strip())
        if e:
            output_formats.append(e)
    return {
        "output_formats":            output_formats or [t.OutputFormat.DXF],
        "autocad_version":           _lookup_enum(t.AcadVersion, d.get("autocad_version")),
        "sheet_size":                _lookup_enum(t.SheetSize, d.get("sheet_size")),
        "drawing_scale":             _lookup_enum(t.DrawingScale, d.get("drawing_scale")),
        "gen_plan_sheet":            _coerce(d.get("gen_plan_sheet"), bool),
        "gen_long_section_sheet":    _coerce(d.get("gen_long_section_sheet"), bool),
        "gen_cross_section_sheet":   _coerce(d.get("gen_cross_section_sheet"), bool),
        "gen_foundation_sheet":      _coerce(d.get("gen_foundation_sheet"), bool),
        "gen_reinforcement_sheet":   _coerce(d.get("gen_reinforcement_sheet"), bool),
        "gen_pier_details_sheet":    _coerce(d.get("gen_pier_details_sheet"), bool),
        "gen_abutment_details_sheet": _coerce(d.get("gen_abutment_details_sheet"), bool),
        "gen_bearings_joints_sheet": _coerce(d.get("gen_bearings_joints_sheet"), bool),
        "gen_boq_sheet":             _coerce(d.get("gen_boq_sheet"), bool),
        "gen_cover_sheet":           _coerce(d.get("gen_cover_sheet"), bool),
        "gen_legends_sheet":         _coerce(d.get("gen_legends_sheet"), bool),
        "gen_hydraulics_sheet":      _coerce(d.get("gen_hydraulics_sheet"), bool),
        "layer_standard":            _lookup_enum(t.LayerStandard, d.get("layer_standard")),
        "title_block_style":         _lookup_enum(t.TitleBlockStyle, d.get("title_block_style")),
        "dimension_precision_mm":    _coerce(d.get("dimension_precision_mm"), int),
        "text_height_scale_factor":  _coerce(d.get("text_height_scale_factor"), Decimal),
        "lineweight_scale_factor":   _coerce(d.get("lineweight_scale_factor"), Decimal),
        "drawing_border_inside_mm":  _coerce(d.get("drawing_border_inside_mm"), int),
        "revision_block_count":      _coerce(d.get("revision_block_count"), int),
        "legend_sheet_required":     _coerce(d.get("legend_sheet_required"), bool),
    }


def _parse_calculations(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "total_length_m":                       _coerce(d.get("total_length_m"), Decimal),
        "span_count":                           _coerce(d.get("span_count"), int),
        "span_depth_ratio":                     _coerce(d.get("span_depth_ratio"), Decimal),
        "deck_area_sqm":                        _coerce(d.get("deck_area_sqm"), Decimal),
        "overall_width_check_ok":               _coerce(d.get("overall_width_check_ok"), bool),
        "concrete_volume_superstructure_cum":   _coerce(d.get("concrete_volume_superstructure_cum"), Decimal),
        "concrete_volume_substructure_cum":     _coerce(d.get("concrete_volume_substructure_cum"), Decimal),
        "concrete_volume_foundation_cum":       _coerce(d.get("concrete_volume_foundation_cum"), Decimal),
        "rebar_weight_tonnes_approx":           _coerce(d.get("rebar_weight_tonnes_approx"), Decimal),
        "prestress_tonnes_approx":              _coerce(d.get("prestress_tonnes_approx"), Decimal),
        "bearing_count":                        _coerce(d.get("bearing_count"), int),
        "expansion_joint_count":                _coerce(d.get("expansion_joint_count"), int),
        "total_estimated_quantity_cost_inr":    _coerce(d.get("total_estimated_quantity_cost_inr"), Decimal),
        "design_scour_vs_foundation_cover_ok":  _coerce(d.get("design_scour_vs_foundation_cover_ok"), bool),
        "waterway_check_pass":                  _coerce(d.get("waterway_check_pass"), bool),
        "afflux_ok":                            _coerce(d.get("afflux_ok"), bool),
        "freeboard_actual_m":                   _coerce(d.get("freeboard_actual_m"), Decimal),
        "minimum_vertical_clearance_m":         _coerce(d.get("minimum_vertical_clearance_m"), Decimal),
        "horizontal_clearance_m":               _coerce(d.get("horizontal_clearance_m"), Decimal),
    }


def _parse_validation_sheet(ws, m, t) -> dict:
    d = _sheet_to_dict(ws)
    return {
        "overall_status":                   _coerce(d.get("overall_status"), str),
        "critical_fail_count":              _coerce(d.get("critical_fail_count"), int),
        "warning_count":                    _coerce(d.get("warning_count"), int),
        "info_count":                       _coerce(d.get("info_count"), int),
        "score_0_to_100":                   _coerce(d.get("score_0_to_100"), Decimal),
        "score_irc05":                      _coerce(d.get("score_irc05"), Decimal),
        "score_irc21":                      _coerce(d.get("score_irc21"), Decimal),
        "score_ircsp55":                    _coerce(d.get("score_ircsp55"), Decimal),
        "score_structural":                 _coerce(d.get("score_structural"), Decimal),
        "score_hydraulic":                  _coerce(d.get("score_hydraulic"), Decimal),
        "score_drawing_standards":          _coerce(d.get("score_drawing_standards"), Decimal),
        "report_html_or_pdf_generated_flag": _coerce(d.get("report_html_or_pdf_generated_flag"), bool),
        "check_version_tag":                _coerce(d.get("check_version_tag"), str),
    }


# ---------------------------------------------------------------------------
# Sheet name normaliser — tolerates minor spelling variations
# ---------------------------------------------------------------------------
_SHEET_ALIASES: dict[str, str] = {
    "project_master": "S01", "project master": "S01", "s01": "S01",
    "bridge_selection": "S02", "bridge selection": "S02", "s02": "S02",
    "geometry_input": "S03", "geometry input": "S03", "s03": "S03",
    "superstructure": "S04", "s04": "S04",
    "substructure": "S05", "s05": "S05",
    "foundation_details": "S06", "foundation details": "S06", "s06": "S06",
    "approaches": "S07", "s07": "S07",
    "hydraulic_data": "S08", "hydraulic data": "S08", "s08": "S08",
    "materials": "S09", "s09": "S09",
    "bearings_joints": "S10", "bearings joints": "S10", "s10": "S10",
    "components_lib": "S11", "components lib": "S11", "s11": "S11",
    "drawing_control": "S12", "drawing control": "S12", "s12": "S12",
    "calculations": "S13", "s13": "S13",
    "validation": "S14", "s14": "S14",
}

_PARSERS = {
    "S01": _parse_project_master,
    "S02": _parse_bridge_selection,
    "S03": _parse_geometry_input,
    "S04": _parse_superstructure,
    "S05": _parse_substructure,
    "S06": _parse_foundation_details,
    "S07": _parse_approaches,
    "S08": _parse_hydraulic_data,
    "S09": _parse_materials,
    "S10": _parse_bearings_joints,
    "S11": _parse_components_lib,
    "S12": _parse_drawing_control,
    "S13": _parse_calculations,
    "S14": _parse_validation_sheet,
}

_MODEL_KEYS = {
    "S01": "project_master",
    "S02": "bridge_selection",
    "S03": "geometry_input",
    "S04": "superstructure",
    "S05": "substructure",
    "S06": "foundation_details",
    "S07": "approaches",
    "S08": "hydraulic_data",
    "S09": "materials",
    "S10": "bearings_joints",
    "S11": "components_lib",
    "S12": "drawing_control",
    "S13": "calculations",
    "S14": "validation",
}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def read_excel(path: str | Path) -> "BridgeProject":  # type: ignore[name-defined]
    """Parse a 14-sheet BridgeCAD Excel workbook and return a BridgeProject.

    Parameters
    ----------
    path : str | Path
        Path to the .xlsx file.

    Returns
    -------
    BridgeProject
        Fully-validated Pydantic model instance.

    Raises
    ------
    ImportError
        If openpyxl is not installed.
    FileNotFoundError
        If the workbook does not exist.
    ValueError
        If the workbook contains no recognisable BridgeCAD sheets.
    """
    if not _OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl is required: pip install openpyxl")

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    m, t = _core()
    wb = openpyxl.load_workbook(path, data_only=True)

    # Map each worksheet to its canonical sheet ID
    sheet_map: dict[str, Any] = {}
    for ws in wb.worksheets:
        key = _SHEET_ALIASES.get(ws.title.strip().lower())
        if key:
            sheet_map[key] = ws

    if not sheet_map:
        raise ValueError(
            f"No recognised BridgeCAD sheets found in '{path.name}'. "
            "Expected sheets named PROJECT_MASTER, BRIDGE_SELECTION, etc."
        )

    # Parse each found sheet; fall back to empty dict for missing ones
    parsed: dict[str, dict] = {}
    for sid, parser in _PARSERS.items():
        ws = sheet_map.get(sid)
        if ws is not None:
            try:
                parsed[sid] = parser(ws, m, t)
            except Exception as exc:
                logger.warning("Error parsing sheet %s: %s", sid, exc)
                parsed[sid] = {}
        else:
            logger.info("Sheet %s not found in workbook — using defaults", sid)
            parsed[sid] = {}

    # Build BridgeProject kwargs, stripping None values so Pydantic defaults apply
    kwargs: dict[str, Any] = {}
    for sid, model_key in _MODEL_KEYS.items():
        sheet_data = {k: v for k, v in parsed[sid].items() if v is not None}
        if sheet_data:
            kwargs[model_key] = sheet_data

    return m.BridgeProject(**kwargs)


__all__ = ["read_excel"]
